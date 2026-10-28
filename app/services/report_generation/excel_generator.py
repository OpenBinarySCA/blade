"""
Excel Generator for BLADE Report

Generates an Excel workbook with multiple sheets containing:
- Summary: Task info, input file info, statistics
- Extraction: Extraction details
- TPLs: Full TPL table
"""

import os
from typing import Dict, Any, List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class ExcelGenerator:
    """Generator for Excel reports with multiple sheets."""

    # BLADE purple color
    PURPLE = "A855F7"
    HEADER_FILL = PatternFill(start_color="E9D5FF", end_color="E9D5FF", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="581C87")
    BORDER = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    def __init__(self, result_data: Dict[str, Any], task_data: Dict[str, Any]):
        """
        Initialize Excel generator.

        Args:
            result_data: The detection result data dictionary
            task_data: Task metadata dictionary
        """
        self.result_data = result_data
        self.task_data = task_data
        self.wb = Workbook()

    def _apply_header_style(self, cell):
        """Apply header styling to a cell."""
        cell.fill = self.HEADER_FILL
        cell.font = self.HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = self.BORDER

    def _apply_cell_style(self, cell):
        """Apply normal cell styling."""
        cell.border = self.BORDER
        cell.alignment = Alignment(vertical='center')

    def _auto_adjust_column_width(self, ws):
        """Auto-adjust column widths based on content."""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)  # Cap at 50
            ws.column_dimensions[column_letter].width = adjusted_width

    def _create_summary_sheet(self):
        """Create Summary sheet with task info, input file info, and statistics."""
        ws = self.wb.active
        ws.title = "Summary"

        row = 1

        # Title
        ws.cell(row=row, column=1, value="BLADE Detection Report - Summary")
        ws.cell(row=row, column=1).font = Font(bold=True, size=16, color=self.PURPLE)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 2

        # Task Information Section
        ws.cell(row=row, column=1, value="Task Information")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12, color=self.PURPLE)
        row += 1

        task_info = [
            ("Task ID", self.task_data.get('task_id', '-')),
            ("Status", self.task_data.get('status', '-')),
            ("Start Time", self.task_data.get('start_at', '-')),
            ("End Time", self.task_data.get('end_at', '-')),
        ]

        for label, value in task_info:
            cell_label = ws.cell(row=row, column=1, value=label)
            cell_value = ws.cell(row=row, column=2, value=str(value) if value else '-')
            self._apply_header_style(cell_label)
            self._apply_cell_style(cell_value)
            row += 1

        row += 1

        # Input File Information Section
        ws.cell(row=row, column=1, value="Input File Information")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12, color=self.PURPLE)
        row += 1

        input_file = self.result_data.get('input_file', {})
        file_info = [
            ("File Name", input_file.get('file_name', '-')),
            ("File Size", input_file.get('file_size', '-')),
            ("MD5", input_file.get('file_md5', '-')),
            ("File Type", input_file.get('readable_type', '-')),
        ]

        for label, value in file_info:
            cell_label = ws.cell(row=row, column=1, value=label)
            cell_value = ws.cell(row=row, column=2, value=str(value) if value else '-')
            self._apply_header_style(cell_label)
            self._apply_cell_style(cell_value)
            row += 1

        row += 1

        # Statistics Section
        ws.cell(row=row, column=1, value="Statistics")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12, color=self.PURPLE)
        row += 1

        extraction = self.result_data.get('extraction', {})
        detection_results = self.result_data.get('detection_results', [])

        # Calculate total libraries
        total_libs = 0
        for result in detection_results:
            if result.get('main_program'):
                total_libs += 1
            total_libs += len(result.get('detected_libraries', []))
            total_libs += len(result.get('detected_dynamic_libraries', []))

        stats = [
            ("Files Extracted", extraction.get('extracted_file_num', 0)),
            ("Binaries Analyzed", len(detection_results)),
            ("Libraries Detected", total_libs),
        ]

        for label, value in stats:
            cell_label = ws.cell(row=row, column=1, value=label)
            cell_value = ws.cell(row=row, column=2, value=value)
            self._apply_header_style(cell_label)
            self._apply_cell_style(cell_value)
            row += 1

        self._auto_adjust_column_width(ws)

    def _create_extraction_sheet(self):
        """Create Extraction sheet with extraction details."""
        ws = self.wb.create_sheet("Extraction")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Extraction Details")
        ws.cell(row=row, column=1).font = Font(bold=True, size=16, color=self.PURPLE)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 2

        extraction = self.result_data.get('extraction', {})

        # Extraction summary
        ws.cell(row=row, column=1, value="Extracted Files Count")
        ws.cell(row=row, column=2, value=extraction.get('extracted_file_num', 0))
        self._apply_header_style(ws.cell(row=row, column=1))
        self._apply_cell_style(ws.cell(row=row, column=2))
        row += 2

        # Analyzed binaries list
        ws.cell(row=row, column=1, value="Analyzed Binaries")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12, color=self.PURPLE)
        row += 1

        # Headers
        headers = ["Index", "Binary Name", "Binary Path", "SHA256"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            self._apply_header_style(cell)
        row += 1

        # Binary list
        detection_results = self.result_data.get('detection_results', [])
        for idx, result in enumerate(detection_results, 1):
            binary_path = (
                result.get('analysis_data', {}).get('target_binary', {}).get('relative_path') or
                result.get('binary_path') or
                'Unknown'
            )
            binary_name = os.path.basename(binary_path)
            sha256 = result.get('binary_sha256', '-')

            cells = [
                ws.cell(row=row, column=1, value=idx),
                ws.cell(row=row, column=2, value=binary_name),
                ws.cell(row=row, column=3, value=binary_path),
                ws.cell(row=row, column=4, value=sha256),
            ]
            for cell in cells:
                self._apply_cell_style(cell)
            row += 1

        self._auto_adjust_column_width(ws)

    def _create_tpls_sheet(self):
        """Create TPLs sheet with full TPL table."""
        ws = self.wb.create_sheet("TPLs")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Detected Third-Party Libraries")
        ws.cell(row=row, column=1).font = Font(bold=True, size=16, color=self.PURPLE)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 2

        # Headers
        headers = ["Index", "Library Name", "Module Name", "Version", "Type", "Binary", "Description"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            self._apply_header_style(cell)
        row += 1

        # Extract and list TPLs
        detection_results = self.result_data.get('detection_results', [])
        idx = 1

        for result in detection_results:
            binary_path = (
                result.get('analysis_data', {}).get('target_binary', {}).get('relative_path') or
                result.get('binary_path') or
                'Unknown'
            )
            binary_name = os.path.basename(binary_path)

            def add_lib_row(lib: Dict[str, Any], lib_type: str):
                nonlocal row, idx
                version = None
                versions = lib.get('versions', [])
                if versions and versions[0].get('version') != 'unknown':
                    version = versions[0].get('version')

                cells = [
                    ws.cell(row=row, column=1, value=idx),
                    ws.cell(row=row, column=2, value=lib.get('name', '-')),
                    ws.cell(row=row, column=3, value=lib.get('component_name') or '-'),
                    ws.cell(row=row, column=4, value=version or '-'),
                    ws.cell(row=row, column=5, value=lib_type),
                    ws.cell(row=row, column=6, value=binary_name),
                    ws.cell(row=row, column=7, value=lib.get('description', '-')),
                ]
                for cell in cells:
                    self._apply_cell_style(cell)
                row += 1
                idx += 1

            # Main program
            main_program = result.get('main_program')
            if main_program and main_program.get('name'):
                add_lib_row(main_program, 'Main Program')

            # Static linked
            for lib in result.get('detected_libraries', []):
                add_lib_row(lib, 'Static Linked')

            # Dynamic linked
            for lib in result.get('detected_dynamic_libraries', []):
                add_lib_row(lib, 'Dynamic Linked')

        self._auto_adjust_column_width(ws)

    def generate(self, output_path: str) -> str:
        """
        Generate Excel workbook with all sheets.

        Args:
            output_path: Path to save the Excel file

        Returns:
            Path to the generated Excel file
        """
        self._create_summary_sheet()
        self._create_extraction_sheet()
        self._create_tpls_sheet()

        self.wb.save(output_path)
        return output_path
